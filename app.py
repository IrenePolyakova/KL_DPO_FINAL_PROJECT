import os
import sys
import streamlit as st
import warnings

# Отключаем предупреждения для torch path
warnings.filterwarnings('ignore', message='.*torch.classes.*')

# Устанавливаем конфигурацию страницы до любых других вызовов Streamlit
st.set_page_config(
    page_title="Оценка качества перевода",
    page_icon="📊",
    layout="wide",
    initial_sidebar_state="expanded"
)

# Импорты остальных модулей
from translation_utils.model_loader import load_model_from_zip
from translation_utils.translator import translate_texts
from translation_utils.metrics import compute_corpus_metrics
from translation_utils.docx_utils import extract_content_from_docx, align_sentences
import pandas as pd
import tempfile
import seaborn as sns
import matplotlib.pyplot as plt
import io
import csv
import codecs
import datetime
from fpdf import FPDF
import zipfile
import re
import traceback
import torch
import gc
import logging
import time
from tqdm import tqdm
from docx import Document
import sacrebleu
from langdetect import detect

# Настройка логирования
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(levelname)s - %(message)s',
    datefmt='%Y-%m-%d %H:%M:%S'
)
logger = logging.getLogger(__name__)

# Инициализация состояния приложения
def init_state():
    """Инициализация состояния приложения"""
    if 'initialized' not in st.session_state:
        st.session_state.initialized = True
        st.session_state.processed_systems = []
        st.session_state.translations = {}
        st.session_state.metrics = {}
        st.session_state.src_sentences = []
        st.session_state.ref_sentences = []
        st.session_state.df = None
        st.session_state.model_speeds = {}
        st.session_state.error_occurred = False
        st.session_state.error_message = ""
        st.session_state.progress = 0
        st.session_state.progress_text = ""

# Инициализация Streamlit
def init_streamlit():
    """Инициализация Streamlit"""
    try:
        # Инициализируем состояние приложения
        init_state()
        
        # Скрываем стандартные стили Streamlit
        hide_streamlit_style = """
            <style>
            #MainMenu {visibility: hidden;}
            footer {visibility: hidden;}
            </style>
        """
        st.markdown(hide_streamlit_style, unsafe_allow_html=True)
        
        return True
    except Exception as e:
        logger.error(f"Ошибка при инициализации Streamlit: {str(e)}")
        return False

# Инициализируем Streamlit при запуске
init_streamlit()

# Отключаем предупреждения Streamlit
st.set_option('client.showErrorDetails', False)
st.set_option('client.toolbarMode', 'minimal')

# Кэширование для оптимизации производительности
@st.cache_resource(ttl=3600)
def load_cached_model(model_zip, model_name):
    """Кэшированная загрузка модели"""
    try:
        logger.info(f"Начало загрузки модели {model_name}")
        with tempfile.TemporaryDirectory() as tmpdir:
            # Сохраняем загруженный файл во временную директорию
            temp_zip_path = os.path.join(tmpdir, f"{model_name}.zip")
            with open(temp_zip_path, "wb") as f:
                f.write(model_zip.getvalue())
            
            model, tokenizer = load_model_from_zip(temp_zip_path, os.path.join(tmpdir, model_name))
            logger.info(f"Модель {model_name} успешно загружена")
            return model, tokenizer
    except Exception as e:
        logger.error(f"Ошибка при загрузке модели {model_name}: {e}")
        logger.error(traceback.format_exc())
        return None, None

@st.cache_data(ttl=3600)
def process_file_content(file, file_type):
    """Кэшированная обработка содержимого файла"""
    if file is None:
        return {'sentences': [], 'references': [], 'tables': []}
    
    try:
        if file_type == 'csv':
            return extract_content_from_csv(file)
        else:
            return extract_content_from_docx(file)
    except Exception as e:
        logger.error(f"Ошибка при обработке файла: {e}")
        return {'sentences': [], 'references': [], 'tables': []}

def clear_gpu_memory():
    """Очистка памяти GPU"""
    if torch.cuda.is_available():
        torch.cuda.empty_cache()
        gc.collect()

# Оптимизированная функция для обработки переводов
def process_model_translation(system_name, model_zip, texts):
    """Обработка переводов для одной модели с загрузкой и управлением памятью"""
    try:
        logger.info(f"Начало обработки переводов для системы {system_name}")
        start_time = time.time()
        
        # Проверяем входные данные
        if not model_zip or not texts:
            logger.error(f"Отсутствуют необходимые данные для {system_name}: model_zip={bool(model_zip)}, texts={bool(texts)}")
            return None
            
        # Загрузка модели с кэшированием
        model, tokenizer = load_cached_model(model_zip, system_name)
        if model is None or tokenizer is None:
            logger.error(f"Не удалось загрузить модель {system_name}")
            return None
            
        # Очищаем память GPU перед переводом
        clear_gpu_memory()
        
        # Определяем оптимальные параметры для данного железа
        device = "cuda" if torch.cuda.is_available() else "cpu"
        
        # Оптимизируем размер батча:
        # - Для GPU (GTX 1050 Ti 4GB): используем меньший батч из-за ограниченной памяти
        # - Для CPU (56 потоков): используем больший батч для лучшей утилизации
        batch_size = 4 if device == "cuda" else 32
        
        # Для CPU используем больше процессов (у нас 56 потоков)
        n_jobs = 1 if device == "cuda" else min(56, len(texts) // 1000 + 1)
        
        # Выполняем перевод с параллельной обработкой
        logger.info(f"Начало перевода {len(texts)} предложений с помощью {system_name}")
        logger.info(f"Параметры: device={device}, batch_size={batch_size}, n_jobs={n_jobs}")
        
        translations, speed = translate_texts(
            texts=texts,
            model=model,
            tokenizer=tokenizer,
            batch_size=batch_size,
            max_length=512,
            num_beams=4,
            device=device,
            show_progress=True,
            n_jobs=n_jobs
        )
        
        # Проверяем результаты перевода
        if not translations:
            logger.error(f"Не получено переводов от модели {system_name}")
            return None
            
        if len(translations) != len(texts):
            logger.error(f"Количество переводов ({len(translations)}) не совпадает с количеством исходных текстов ({len(texts)})")
            return None
        
        # Очистка памяти после перевода
        logger.info("Очистка памяти после перевода")
        del model
        del tokenizer
        torch.cuda.empty_cache()
        gc.collect()
        
        logger.info(f"Перевод завершен за {time.time() - start_time:.2f} сек, скорость: {speed:.2f} предл./сек")
        return translations
        
    except Exception as e:
        logger.error(f"Ошибка при обработке переводов для {system_name}: {str(e)}")
        logger.error(traceback.format_exc())
        return None

# Функция для генерации PDF

def generate_pdf_report(df, metrics_df, fig):
    pdf = FPDF()
    pdf.add_page()
    # Используем стандартный шрифт 'helvetica'
    pdf.set_font("helvetica", size=12)
    
    # Заголовок
    pdf.set_font("helvetica", 'B', 16)
    pdf.cell(200, 10, txt="Отчет о сравнении систем перевода", ln=1, align='C')
    pdf.set_font("helvetica", size=12)
    pdf.cell(200, 10, txt=f"Дата: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}", ln=1)
    pdf.ln(10)
    
    # Метрики по корпусу
    if metrics_df is not None:
        pdf.set_font("helvetica", 'B', 14)
        pdf.cell(200, 10, txt="Метрики по корпусу", ln=1)
        pdf.set_font("helvetica", size=10)
        
        # Заголовки таблицы
        col_width = 40
        pdf.cell(col_width, 10, "Система", border=1)
        pdf.cell(col_width, 10, "BLEU", border=1)
        pdf.cell(col_width, 10, "chrF", border=1)
        pdf.cell(col_width, 10, "TER", border=1)
        pdf.ln()
        
        # Данные
        for _, row in metrics_df.iterrows():
            pdf.cell(col_width, 10, str(row['Система']), border=1)
            pdf.cell(col_width, 10, f"{row['BLEU']:.2f}", border=1)
            pdf.cell(col_width, 10, f"{row['chrF']:.2f}", border=1)
            pdf.cell(col_width, 10, f"{row['TER']:.2f}", border=1)
            pdf.ln()
        pdf.ln(10)
    
    # График
    if fig is not None:
        pdf.set_font("helvetica", 'B', 14)
        pdf.cell(200, 10, txt="Визуализация метрик", ln=1)
        
        # Сохраняем график в буфер памяти
        img_buffer = io.BytesIO()
        fig.savefig(img_buffer, format='png', bbox_inches='tight')
        img_buffer.seek(0)
        
        # Используем временный файл
        with tempfile.NamedTemporaryFile(suffix='.png', delete=True) as tmpfile:
            tmpfile.write(img_buffer.getvalue())
            tmpfile.flush()
            pdf.image(tmpfile.name, x=10, y=pdf.get_y(), w=180)
        
        pdf.ln(90)
    
    # Пояснения к метрикам
    pdf.set_font("helvetica", 'B', 14)
    pdf.cell(200, 10, txt="Пояснения к метрикам", ln=1)
    pdf.set_font("helvetica", size=10)
    explanations = [
        "BLEU (Bilingual Evaluation Understudy) - метрика, основанная на точности n-грамм. "
        "Учитывает совпадение последовательностей слов между переводом и эталоном. "
        "Значение от 0 до 100, чем выше, тем лучше.",
        
        "chrF (Character n-gram F-score) - метрика, основанная на F-мере для символьных n-грамм (обычно n=6). "
        "Учитывает точность и полноту на уровне символов. "
        "Значение от 0 до 1 (или в процентах), чем выше, тем лучше.",
        
        "TER (Translation Edit Rate) - метрика, измеряющая минимальное количество редактирований "
        "(вставка, удаление, замена, перестановка слов), необходимых для преобразования перевода в эталон, "
        "нормализованное на длину эталона. Чем ниже значение, тем лучше."
    ]
    
    for exp in explanations:
        pdf.multi_cell(0, 5, txt=exp)
        pdf.ln(3)
    
    return pdf.output(dest='S')

def create_zip_archive(src_sentences, ref_sentences, systems, tables_systems, df, metrics_df, fig, file_prefix):
    """Создает ZIP-архив со всеми результатами"""
    zip_buffer = io.BytesIO()
    
    with zipfile.ZipFile(zip_buffer, 'w', zipfile.ZIP_DEFLATED) as zip_file:
        # Исходные данные
        zip_file.writestr(f"source_{file_prefix}.txt", "\n".join(src_sentences))
        
        if ref_sentences:
            zip_file.writestr(f"reference_{file_prefix}.txt", "\n".join(ref_sentences))
        
        # Переводы предложений
        for system in systems:
            if system in st.session_state.translations:
                zip_file.writestr(f"translation_{system}_{file_prefix}.txt", 
                                 "\n".join(st.session_state.translations[system]))
        
        # Сравнение
        csv_data = df.to_csv(index=False).encode('utf-8-sig')
        zip_file.writestr(f"comparison_{file_prefix}.csv", csv_data)
        
        # Метрики
        if metrics_df is not None:
            metrics_csv = metrics_df.to_csv(index=False).encode('utf-8-sig')
            zip_file.writestr(f"metrics_{file_prefix}.csv", metrics_csv)
            
            # График
            if fig:
                img_io = io.BytesIO()
                fig.savefig(img_io, format='png', bbox_inches='tight')
                img_bytes = img_io.getvalue()
                zip_file.writestr(f"metrics_plot_{file_prefix}.png", img_bytes)
            
            # PDF отчет
            try:
                pdf_bytes = generate_pdf_report(df, metrics_df, fig)
                zip_file.writestr(f"report_{file_prefix}.pdf", pdf_bytes)
            except Exception:
                pass
    
    zip_buffer.seek(0)
    return zip_buffer

# Функция для извлечения контента из CSV
def extract_content_from_csv(file):
    """Извлекает предложения и переводы из CSV файла"""
    if file is None:
        return {'sentences': [], 'references': [], 'source_translations': [], 'tables': [], 'dataframe': None}
    
    try:
        file.seek(0)
        content = file.read()
        decoded = content.decode('utf-8-sig') if isinstance(content, bytes) else content
        
        df = pd.read_csv(io.StringIO(decoded))
        file_name = os.path.basename(file.name) if hasattr(file, 'name') else 'unknown'
        
        # Определяем колонки
        columns_lower = {col.lower(): col for col in df.columns}
        translations = {}
        
        # Ищем колонки с английским и русским текстом
        for col in df.columns:
            col_lower = col.lower()
            # Английский текст
            if any(x in col_lower for x in ['source', 'src', 'en', 'исходный', 'english']):
                translations['source'] = df[col].fillna('').tolist()
            # Русский текст/переводы
            elif any(x in col_lower for x in ['target', 'tgt', 'ru', 'эталон', 'перевод', 'russian']):
                translations[f'ru_from_{file_name}'] = df[col].fillna('').tolist()

        if not translations:
            # Если не нашли по ключевым словам, берем первые две колонки
            if len(df.columns) >= 2:
                translations['source'] = df[df.columns[0]].fillna('').tolist()
                translations[f'ru_from_{file_name}'] = df[df.columns[1]].fillna('').tolist()
            else:
                translations['source'] = df[df.columns[0]].fillna('').tolist()

        # Создаем DataFrame для отображения
        new_df = pd.DataFrame()
        new_df['ID'] = range(1, len(df) + 1)
        
        if 'source' in translations:
            new_df[f"Source ({file_name})"] = translations['source']
        
        # Добавляем все найденные русские переводы
        for key, value in translations.items():
            if key.startswith('ru_from_'):
                new_df[f"Translation from {file_name}"] = value

        return {
            'sentences': translations.get('source', []),
            'references': translations.get(f'ru_from_{file_name}', []),
            'source_translations': translations.get(f'ru_from_{file_name}', []),
            'tables': [],
            'dataframe': new_df,
            'translations': translations
        }

    except Exception as e:
        logger.error(f"Ошибка при чтении CSV: {str(e)}")
        return {'sentences': [], 'references': [], 'source_translations': [], 'tables': [], 'dataframe': None}

def export_tmx(src_sentences, tgt_sentences, system_name):
    """Экспортирует пары предложений в TMX-формат"""
    from translation_utils.tmx_utils import create_tmx_file
    import io
    tmx_io = io.BytesIO()
    create_tmx_file(
        src_sentences=src_sentences,
        tgt_sentences=tgt_sentences,
        fileobj=tmx_io,
        sys_name=system_name
    )
    tmx_io.seek(0)
    return tmx_io.read()

def display_results(
    df, systems, source_file=None, ref_file=None, model1_file=None, model2_file=None,
    yandex_file=None, google_file=None, deepl_file=None, extra_files=None, extra_names=None,
    model1_name=None, model2_name=None
):
    try:
        results_df = pd.DataFrame()
        
        if isinstance(df, dict):
            # Добавляем ID
            num_rows = len(df.get('sentences', []))
            results_df['ID'] = range(1, num_rows + 1)
            
            # Добавляем исходный текст
            source_name = f"Source ({os.path.basename(source_file.name)})" if source_file else "Source"
            results_df[source_name] = df.get('sentences', [])
            
            # Добавляем перевод из исходного файла, если есть
            if df.get('source_translations'):
                source_trans_name = f"Translation from {os.path.basename(source_file.name)}"
                results_df[source_trans_name] = df['source_translations']
            
            # Добавляем перевод из эталонного файла
            if ref_file and df.get('references'):
                ref_name = f"Translation from {os.path.basename(ref_file.name)}"
                results_df[ref_name] = df['references']
            
            # Добавляем переводы от систем машинного перевода
            for system, file in [
                ('yandex', yandex_file),
                ('google', google_file),
                ('deepl', deepl_file)
            ]:
                if file:
                    file_name = os.path.basename(file.name)
                    try:
                        content = extract_content_from_csv(file) if file.name.endswith('.csv') else extract_content_from_docx(file)
                        if content.get('references'):
                            results_df[f"Translation from {file_name}"] = content['references']
                    except Exception as e:
                        logger.error(f"Ошибка при обработке файла {file_name}: {str(e)}")
            
            # Добавляем переводы от моделей
            if 'translations' in df:
                for system in systems:
                    if system in df['translations']:
                        file_name = None
                        if system == model1_name and model1_file:
                            file_name = os.path.basename(model1_file.name)
                        elif system == model2_name and model2_file:
                            file_name = os.path.basename(model2_file.name)
                        
                        system_name = f"Translation {system} ({file_name})" if file_name else f"Translation ({system})"
                        results_df[system_name] = df['translations'][system]
            
            # Добавляем дополнительные переводы
            if extra_files and extra_names:
                for file, name in zip(extra_files, extra_names):
                    if file:
                        try:
                            content = extract_content_from_csv(file) if file.name.endswith('.csv') else extract_content_from_docx(file)
                            if content.get('references'):
                                results_df[f"Translation from {os.path.basename(file.name)}"] = content['references']
                        except Exception as e:
                            logger.error(f"Ошибка при обработке дополнительного файла {file.name}: {str(e)}")
        
        elif isinstance(df, pd.DataFrame):
            results_df = df.copy()
        else:
            raise ValueError("Неподдерживаемый тип входных данных")

        if results_df.empty:
            st.warning("Нет данных для отображения")
            return None

        st.dataframe(results_df, use_container_width=True)
        return results_df

    except Exception as e:
        logger.error(f"Ошибка при отображении результатов: {str(e)}")
        st.error(f"❌ Произошла ошибка при отображении результатов: {str(e)}")
        return None

# Основной интерфейс
def main():
    # Инициализируем Streamlit
    if not init_streamlit():
        st.error("Не удалось инициализировать приложение. Пожалуйста, проверьте логи для получения дополнительной информации.")
        return
    
    st.title("📄 Сравнение переводов моделей и машинных переводчиков")
    current_env = os.environ.get("CONDA_DEFAULT_ENV", "неизвестно")
    python_version = f"{sys.version_info.major}.{sys.version_info.minor}.{sys.version_info.micro}"
    st.markdown(f"""
    **Conda-окружение:** `{current_env}`  
    **Версия Python:** `{python_version}`
    """)
    st.divider()
    
    col1, col2 = st.columns(2)
    with col1:
        uploaded_source = st.file_uploader("📘 Исходный документ (обязательно)", type=["docx", "csv"])
        model1_zip = st.file_uploader("🤖 Дообученная модель (zip)", type=["zip"])
        model2_zip = st.file_uploader("🧑‍💻 Исходная модель (zip, опционально)", type=["zip"])
    with col2:
        uploaded_ref = st.file_uploader("📗 Эталонный перевод (опционально)", type=["docx", "csv"])
        yandex_file = st.file_uploader("🌐 Перевод Яндекс (docx/csv, опционально)", type=["docx", "csv"])
        google_file = st.file_uploader("🌐 Перевод Google (docx/csv, опционально)", type=["docx", "csv"])
        deepl_file = st.file_uploader("🌐 Перевод DeepL (docx/csv, опционально)", type=["docx", "csv"])
        extra_count = st.number_input("Сколько дополнительных переводов загрузить?", min_value=0, max_value=5, value=0, step=1)
        extra_files = []
        extra_names = []
        for i in range(extra_count):
            extra_files.append(st.file_uploader(f"Доп. перевод {i+1} (docx/csv)", type=["docx", "csv"], key=f"extra_file_{i}"))
            extra_names.append(st.text_input(f"Название системы для доп. перевода {i+1}", key=f"extra_name_{i}"))

    # Названия моделей
    col1, col2 = st.columns(2)
    with col1:
        model1_name = st.text_input("Название Модели 1", value="Дообученная модель")
    with col2:
        model2_name = st.text_input("Название Модели 2", value="Исходная модель")

    if st.button("🚀 Сравнить переводы"):
        if not uploaded_source or not model1_zip:
            st.warning("Пожалуйста, загрузите исходный документ и дообученную модель.")
            return
        
        progress_bar = st.progress(0)
        status_text = st.empty()
        
        def update_progress(step, total, msg):
            progress = int((step/total)*100)
            progress_bar.progress(progress)
            status_text.text(f"{msg} ({step}/{total})")
            st.session_state.progress = progress
            st.session_state.progress_text = msg
        
        try:
            # Сброс состояния ошибок
            st.session_state.error_occurred = False
            st.session_state.error_message = ""
            
            # Обработка исходного файла
            update_progress(1, 10, "Обработка исходного документа")
            if uploaded_source.name.endswith('.csv'):
                src_content = extract_content_from_csv(uploaded_source)
                src_sentences = src_content['sentences']  # Берем английский текст
                if not uploaded_ref:  # Если эталонный перевод не загружен отдельно, берем его из исходного файла
                    ref_sentences = src_content['references']  # Берем русский перевод
            else:
                src_content = extract_content_from_docx(uploaded_source)
                src_sentences = src_content['sentences']
            
            # Обработка эталонного перевода, если загружен отдельно
            if uploaded_ref:
                update_progress(2, 10, "Обработка эталонного перевода")
                if uploaded_ref.name.endswith('.csv'):
                    ref_content = extract_content_from_csv(uploaded_ref)
                    ref_sentences = ref_content['references']  # Берем русский перевод
                else:
                    ref_content = extract_content_from_docx(uploaded_ref)
                    ref_sentences = ref_content['sentences']
            
            # Сохраняем предложения в состояние
            st.session_state.src_sentences = src_sentences
            st.session_state.ref_sentences = ref_sentences
            
            # Обработка моделей
            update_progress(3, 10, "Загрузка моделей")
            with tempfile.TemporaryDirectory() as tmpdir:
                # Обработка первой модели
                model1_preds = process_model_translation(model1_name, model1_zip, src_sentences)
                if model1_preds:
                    st.session_state.translations[model1_name] = model1_preds
                    st.session_state.processed_systems.append(model1_name)
                    update_progress(5, 10, f"Перевод {model1_name} завершен")
                else:
                    st.error(f"Ошибка при обработке модели {model1_name}")
                    st.session_state.error_occurred = True
                    st.session_state.error_message = f"Ошибка при обработке модели {model1_name}"
                
                # Обработка второй модели
                if model2_zip:
                    model2_preds = process_model_translation(model2_name, model2_zip, src_sentences)
                    if model2_preds:
                        st.session_state.translations[model2_name] = model2_preds
                        st.session_state.processed_systems.append(model2_name)
                        update_progress(6, 10, f"Перевод {model2_name} завершен")
                    else:
                        st.error(f"Ошибка при обработке модели {model2_name}")
                
                # Обработка переводов Yandex, Google, DeepL
                update_progress(7, 10, "Обработка машинных переводов")
                if yandex_file:
                    if yandex_file.name.endswith('.csv'):
                        yandex_content = extract_content_from_csv(yandex_file)
                        yandex_sentences = yandex_content['references']  # Берем русский перевод из колонки 'ru'
                    else:
                        yandex_content = extract_content_from_docx(yandex_file)
                        yandex_sentences = yandex_content['sentences']
                    st.session_state.translations['yandex'] = yandex_sentences
                    st.session_state.processed_systems.append('yandex')
                
                if google_file:
                    if google_file.name.endswith('.csv'):
                        google_content = extract_content_from_csv(google_file)
                        google_sentences = google_content['references']  # Берем русский перевод из колонки 'ru'
                    else:
                        google_content = extract_content_from_docx(google_file)
                        google_sentences = google_content['sentences']
                    st.session_state.translations['google'] = google_sentences
                    st.session_state.processed_systems.append('google')
                
                if deepl_file:
                    if deepl_file.name.endswith('.csv'):
                        deepl_content = extract_content_from_csv(deepl_file)
                        deepl_sentences = deepl_content['references']  # Берем русский перевод из колонки 'ru'
                    else:
                        deepl_content = extract_content_from_docx(deepl_file)
                        deepl_sentences = deepl_content['sentences']
                    st.session_state.translations['deepl'] = deepl_sentences
                    st.session_state.processed_systems.append('deepl')
                
                # Обработка дополнительных переводов
                for i, (extra_file, extra_name) in enumerate(zip(extra_files, extra_names)):
                    if extra_file:
                        update_progress(7+i, 10, f"Обработка {extra_name}")
                        if extra_file.name.endswith('.csv'):
                            extra_content = extract_content_from_csv(extra_file)
                            extra_sentences = extra_content['references']  # Берем русский перевод из колонки 'ru'
                        else:
                            extra_content = extract_content_from_docx(extra_file)
                            extra_sentences = extra_content['sentences']
                        st.session_state.translations[extra_name] = extra_sentences
                        st.session_state.processed_systems.append(extra_name)
            
            update_progress(9, 10, "Формирование результатов")
            # Создаем DataFrame с результатами
            data = []
            for i, src in enumerate(st.session_state.src_sentences):
                row = {'ID': i+1, 'Исходный текст': src}
                if st.session_state.ref_sentences and i < len(st.session_state.ref_sentences):
                    row['Эталонный перевод'] = st.session_state.ref_sentences[i]
                
                for system in st.session_state.processed_systems:
                    if system in st.session_state.translations and i < len(st.session_state.translations[system]):
                        row[f'Перевод {system}'] = st.session_state.translations[system][i]
                
                data.append(row)
            
            df = pd.DataFrame(data)
            st.session_state.df = df
            
            # ---- КОРПУСНЫЕ МЕТРИКИ ----
            metrics_df = None
            fig = None
            if st.session_state.ref_sentences:
                metrics_data = []
                for system in st.session_state.processed_systems:
                    try:
                        preds = st.session_state.translations[system]
                        
                        # Проверяем соответствие длин
                        if len(preds) != len(st.session_state.ref_sentences):
                            logger.warning(
                                f"Разное количество предложений для системы {system}: "
                                f"переводы={len(preds)}, эталон={len(st.session_state.ref_sentences)}"
                            )
                            continue
                        
                        # Фильтруем пустые строки
                        valid_pairs = [
                            (pred, ref) for pred, ref in zip(preds, st.session_state.ref_sentences)
                            if pred and pred.strip() and ref and ref.strip()
                        ]
                        
                        if not valid_pairs:
                            logger.warning(f"Нет валидных пар перевод-эталон для системы {system}")
                            continue
                            
                        filtered_preds, filtered_refs = zip(*valid_pairs)
                        
                        # Вычисляем метрики
                        bleu = sacrebleu.corpus_bleu(filtered_preds, [filtered_refs]).score
                        chrf = sacrebleu.corpus_chrf(filtered_preds, [filtered_refs]).score
                        ter = sacrebleu.corpus_ter(filtered_preds, [filtered_refs]).score
                        
                        metrics_data.append({
                            'Система': system,
                            'BLEU': round(bleu,2),
                            'chrF': round(chrf,2),
                            'TER': round(ter,2)
                        })
                        
                    except Exception as e:
                        logger.error(f"Ошибка при вычислении метрик для системы {system}: {str(e)}")
                        continue
                
                if metrics_data:
                    metrics_df = pd.DataFrame(metrics_data)
                    # Проверка наличия нужных колонок
                    if all(col in metrics_df.columns for col in ['Система', 'BLEU', 'chrF', 'TER']):
                        fig, ax = plt.subplots(figsize=(10, 6))
                        metrics_plot = metrics_df.melt(
                            id_vars=['Система'],
                            value_vars=['BLEU','chrF','TER'],
                            var_name='Метрика',
                            value_name='Значение'
                        )
                        sns.barplot(data=metrics_plot, x='Система', y='Значение', hue='Метрика', ax=ax)
                        ax.set_title('Корпусные метрики по системам')
                        ax.set_ylabel('Значение')
                        ax.set_xlabel('Система')
                        ax.legend(title='Метрика')
                        plt.tight_layout()
                        st.session_state.fig = fig
                    else:
                        st.warning("Некорректный формат данных метрик.")
                        st.session_state.fig = None
                else:
                    st.warning("Нет данных для построения метрик.")
                    st.session_state.fig = None

            st.session_state.metrics_df = metrics_df
            st.session_state.fig = fig
            
            update_progress(10, 10, "Обработка завершена")
            st.success("✅ Обработка завершена!")
            
        except Exception as e:
            st.session_state.error_occurred = True
            st.session_state.error_message = str(e)
            st.error(f"❌ Ошибка при обработке файлов: {str(e)}")
            st.text(traceback.format_exc())
        
        # Отображение результатов
        if not st.session_state.error_occurred:
            display_results(
                st.session_state.df, 
                st.session_state.processed_systems,
                source_file=uploaded_source,
                ref_file=uploaded_ref,
                model1_file=model1_zip,
                model2_file=model2_zip,
                yandex_file=yandex_file,
                google_file=google_file,
                deepl_file=deepl_file,
                extra_files=extra_files,
                extra_names=extra_names,
                model1_name=model1_name,
                model2_name=model2_name
            )
            
            # Отображение метрик
            if st.session_state.metrics_df is not None:
                st.subheader("📈 Метрики по корпусу")
                st.dataframe(st.session_state.metrics_df, use_container_width=True)
                
                st.subheader("📊 Визуализация метрик")
                st.pyplot(st.session_state.fig)
            
            st.subheader("⬇️ Экспорт результатов")
            
            # TMX для каждой системы
            tmx_col1, tmx_col2 = st.columns(2)
            for idx, system in enumerate(st.session_state.processed_systems):
                if system in st.session_state.translations:
                    tmx_bytes = export_tmx(
                        st.session_state.src_sentences, 
                        st.session_state.translations[system], 
                        system
                    )
                    # Чередуем колонки для кнопок
                    with tmx_col1 if idx % 2 == 0 else tmx_col2:
                        st.download_button(
                            f"TMX для {system}", 
                            tmx_bytes, 
                            file_name=f"{system}.tmx", 
                            mime="application/octet-stream",
                            key=f"tmx_download_{system}_{idx}"  # Добавляем индекс для уникальности
                        )
            
            # ZIP архив
            if uploaded_source:
                dt = datetime.datetime.now().strftime("%Y-%m-%dT%H-%M-%S")
                src_folder = os.path.splitext(os.path.basename(uploaded_source.name))[0]
                zip_name = f"{dt}_{src_folder}.zip"
                
                zip_buffer = create_zip_archive(
                    st.session_state.src_sentences,
                    st.session_state.ref_sentences,
                    st.session_state.processed_systems,
                    [],
                    st.session_state.df,
                    st.session_state.metrics_df,
                    st.session_state.fig,
                    f"{dt}_{src_folder}"
                )
                
                st.download_button(
                    "Скачать все результаты (ZIP)", 
                    zip_buffer, 
                    file_name=zip_name, 
                    mime="application/zip",
                    key=f"zip_download_{dt}"  # Используем timestamp для уникальности
                )

# Основной блок выполнения
if __name__ == "__main__":
    main()

# --- Конец файла ---